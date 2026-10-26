#!/usr/bin/env python3
"""Backlink/outreach summary from outreach_tracker.xlsx for the weekly SEO report.
Writes reports/outreach-summary.json (and prev-outreach-summary.json is kept by the caller)."""
import json, collections, datetime, pathlib, re
import openpyxl

root = pathlib.Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(root / 'outreach_tracker.xlsx', data_only=True)
rows = list(wb['Outreach Tracker'].iter_rows(min_row=2, values_only=True))
rows = [r for r in rows if r[1]]
status = lambda s: (s or '').lower()
def bucket(s):
    s = status(s)
    if s.startswith('bounced'): return 'bounced'
    if 'replied' in s or s.startswith('replied'): return 'replied'
    if s.startswith('sent'): return 'sent'
    if s.startswith('declined') or 'do not contact' in s: return 'declined'
    if s.startswith(('skipped', 'researched', 'draft', 'actionable')): return 'not_sent'
    return 'other'
buckets = collections.Counter(bucket(r[6]) for r in rows)
resp = collections.Counter(r[10] for r in rows if r[10])
dates = collections.Counter()
for r in rows:
    d = r[7]
    if isinstance(d, (datetime.date, datetime.datetime)):
        dates[d.strftime('%Y-%m-%d')] += 1
live = [r[1] for r in rows if re.search(r'link (live|placed|published)|published', (r[9] or '') + (r[6] or ''), re.I)]
out = {
    'prospects_total': len(rows),
    'status_buckets': dict(buckets),
    'response_types': dict(resp),
    'sent_by_date': dict(sorted(dates.items())[-14:]),
    'possible_live_links': live,
    'responses_sheet_rows': wb['Responses'].max_row - 1,
    'competitors_researched': wb['Competitors Researched'].max_row - 1,
    'triage_rows': wb['Full Backlink Triage'].max_row - 1,
}
(root / 'reports' / 'outreach-summary.json').write_text(json.dumps(out, indent=2))
print(json.dumps(out, indent=2))
